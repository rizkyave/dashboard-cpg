import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import shutil
import os, sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from scripts.classify_rules import classify_record

SRC_FILE = 'Daftar_Barang_Gabungan_Semua_Perusahaan.xlsx'
BAK_FILE = 'Daftar_Barang_Gabungan_Semua_Perusahaan_ORIGINAL.xlsx'
OUT_FILE = 'Daftar_Barang_Gabungan_Semua_Perusahaan_Enriched.xlsx'

# Always reload from the pure pristine backup
if not os.path.exists(BAK_FILE):
    shutil.copyfile(SRC_FILE, BAK_FILE)
    print(f'Backup created: {BAK_FILE}')

wb = openpyxl.load_workbook(BAK_FILE)

header_fill = PatternFill(start_color='1F4E79', end_color='1F4E79', fill_type='solid') # Navy
new_col_fill = PatternFill(start_color='2F5597', end_color='2F5597', fill_type='solid') # Royal Navy
header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
regular_font = Font(name='Calibri', size=10)
bold_font = Font(name='Calibri', size=10, bold=True)

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

align_center = Alignment(horizontal='center', vertical='center')
align_left = Alignment(horizontal='left', vertical='center')
align_right = Alignment(horizontal='right', vertical='center')

new_headers = [
    'Kategori Utama',
    'Jenis / Sub-Kategori',
    'Satuan (UoM)',
    'Sifat Entitas',
    'Status Stok Logistik',
    'Rekomendasi Tindakan Logistik'
]

# Process Data Sheets
sheet_names = ['Gabungan Semua Perusahaan', 'PT MANDAR OCEAN', 'PT HANA LINES', 'PT CINDARA PRATAMA LINES']

stats_per_cat = {}
stats_per_sifat = {}

for sname in sheet_names:
    ws = wb[sname]
    is_gabungan = (sname == 'Gabungan Semua Perusahaan')
    
    start_col = 9 if is_gabungan else 8
    for idx, nh in enumerate(new_headers, start=start_col):
        cell = ws.cell(row=1, column=idx, value=nh)
        cell.fill = new_col_fill
        cell.font = header_font
        cell.alignment = align_center

    max_row = ws.max_row
    print(f'Processing {sname} ({max_row - 1} rows)...')
    
    for r in range(2, max_row + 1):
        if is_gabungan:
            pt = ws.cell(row=r, column=1).value
            no_brg = ws.cell(row=r, column=2).value
            desc = ws.cell(row=r, column=3).value
            qty = ws.cell(row=r, column=4).value or 0
            tb = ws.cell(row=r, column=6).value
        else:
            pt = sname
            no_brg = ws.cell(row=r, column=1).value
            desc = ws.cell(row=r, column=2).value
            qty = ws.cell(row=r, column=3).value or 0
            tb = ws.cell(row=r, column=5).value
            
        cat, sub, uom, sifat, act = classify_record(no_brg, desc, qty, tb)
        
        # Track stats for Gabungan
        if is_gabungan:
            if cat not in stats_per_cat:
                stats_per_cat[cat] = {'total_item': 0, 'total_qty': 0, 'stok_aktif': 0}
            stats_per_cat[cat]['total_item'] += 1
            stats_per_cat[cat]['total_qty'] += qty
            if qty > 0:
                stats_per_cat[cat]['stok_aktif'] += 1
                
            if sifat not in stats_per_sifat:
                stats_per_sifat[sifat] = {'total_item': 0, 'total_qty': 0}
            stats_per_sifat[sifat]['total_item'] += 1
            stats_per_sifat[sifat]['total_qty'] += qty

        # Status Stok
        if qty > 0:
            stok_status = 'Stok Aktif (>0)'
        elif qty < 0:
            stok_status = 'Stok Minus (<0)'
        else:
            stok_status = 'Stok Kosong (0)'
            
        row_enrichment = [cat, sub, uom, sifat, stok_status, act]
        
        for c_idx, val in enumerate(row_enrichment, start=start_col):
            c_cell = ws.cell(row=r, column=c_idx, value=val)
            c_cell.font = regular_font
            c_cell.border = thin_border
            if c_idx in [start_col + 2, start_col + 3, start_col + 4]: # UoM, Sifat, Status
                c_cell.alignment = align_center
            else:
                c_cell.alignment = align_left
                
            if val == 'Stok Minus (<0)':
                c_cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')
                c_cell.font = Font(name='Calibri', size=10, bold=True, color='9C0006')
            elif val == 'Keluarkan dari Master Persediaan':
                c_cell.fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')

    for col in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col)
        h_val = ws.cell(row=1, column=col).value or ''
        ws.column_dimensions[col_letter].width = max(len(str(h_val)) + 4, 15)

print('Data sheets enriched successfully.')

# Update Dashboard / Ringkasan Sheet
if 'Ringkasan & Dashboard' in wb.sheetnames:
    ws_dash = wb['Ringkasan & Dashboard']
    
    # 3. Enrich Table 3 (Top 15) with UoM, Kategori Utama, Sifat Entitas
    t3_add_headers = ['Satuan (UoM)', 'Kategori Utama', 'Sifat Entitas']
    for idx, th in enumerate(t3_add_headers, start=8):
        c = ws_dash.cell(row=24, column=idx, value=th)
        c.fill = header_fill
        c.font = header_font
        c.alignment = align_center
        
    for r in range(25, 40): # Top 15 rows
        no_brg = ws_dash.cell(row=r, column=3).value
        desc = ws_dash.cell(row=r, column=4).value
        qty = ws_dash.cell(row=r, column=5).value or 0
        tb = ws_dash.cell(row=r, column=6).value
        
        cat, sub, uom, sifat, act = classify_record(no_brg, desc, qty, tb)
        
        c_uom = ws_dash.cell(row=r, column=8, value=uom)
        c_uom.alignment = align_center
        c_uom.font = regular_font
        c_uom.border = thin_border
        
        c_cat = ws_dash.cell(row=r, column=9, value=cat)
        c_cat.alignment = align_left
        c_cat.font = regular_font
        c_cat.border = thin_border
        
        c_sifat = ws_dash.cell(row=r, column=10, value=sifat)
        c_sifat.alignment = align_center
        c_sifat.font = regular_font
        c_sifat.border = thin_border
        if sifat != 'Barang Fisik':
            c_sifat.fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')
            c_sifat.font = Font(name='Calibri', size=10, bold=True, color='B25900')

    # Table 4 starts at row 42
    r_start = 42
    ws_dash.cell(row=r_start, column=1, value='4. Rekapitulasi Berdasarkan Kategori Utama Persediaan (Konsolidasi)').font = Font(name='Calibri', size=12, bold=True, color='1F4E79')
    r_start += 1
    
    dash_h4 = ['No', 'Kategori Utama', 'Total Item', 'Item Stok Aktif (>0)', 'Total Kuantitas Fisik', '% Portofolio Item']
    for idx, dh in enumerate(dash_h4, start=1):
        cell = ws_dash.cell(row=r_start, column=idx, value=dh)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
    
    r_start += 1
    sorted_cats = sorted(stats_per_cat.items(), key=lambda x: x[1]['total_item'], reverse=True)
    
    grand_items = sum(v['total_item'] for v in stats_per_cat.values())
    grand_qty = sum(v['total_qty'] for v in stats_per_cat.values())
    grand_aktif = sum(v['stok_aktif'] for v in stats_per_cat.values())
    
    row_no = 1
    for cat_name, data in sorted_cats:
        ws_dash.cell(row=r_start, column=1, value=row_no).alignment = align_center
        ws_dash.cell(row=r_start, column=2, value=cat_name).alignment = align_left
        
        c3 = ws_dash.cell(row=r_start, column=3, value=data['total_item'])
        c3.alignment = align_right
        c3.number_format = '#,##0'
        
        c4 = ws_dash.cell(row=r_start, column=4, value=data['stok_aktif'])
        c4.alignment = align_right
        c4.number_format = '#,##0'
        
        c5 = ws_dash.cell(row=r_start, column=5, value=data['total_qty'])
        c5.alignment = align_right
        c5.number_format = '#,##0'
        
        pct_val = f"{data['total_item']/grand_items*100:.1f}%"
        ws_dash.cell(row=r_start, column=6, value=pct_val).alignment = align_center
        
        for c in range(1, 7):
            ws_dash.cell(row=r_start, column=c).border = thin_border
            ws_dash.cell(row=r_start, column=c).font = regular_font
        
        r_start += 1
        row_no += 1
        
    # Total row for Table 4
    ws_dash.cell(row=r_start, column=1, value='').alignment = align_center
    ws_dash.cell(row=r_start, column=2, value='TOTAL KONSOLIDASI').font = bold_font
    
    c_tot3 = ws_dash.cell(row=r_start, column=3, value=grand_items)
    c_tot3.font = bold_font
    c_tot3.alignment = align_right
    c_tot3.number_format = '#,##0'
    
    c_tot4 = ws_dash.cell(row=r_start, column=4, value=grand_aktif)
    c_tot4.font = bold_font
    c_tot4.alignment = align_right
    c_tot4.number_format = '#,##0'
    
    c_tot5 = ws_dash.cell(row=r_start, column=5, value=grand_qty)
    c_tot5.font = bold_font
    c_tot5.alignment = align_right
    c_tot5.number_format = '#,##0'
    
    ws_dash.cell(row=r_start, column=6, value='100.0%').font = bold_font
    ws_dash.cell(row=r_start, column=6).alignment = align_center
    for c in range(1, 7):
        ws_dash.cell(row=r_start, column=c).border = thin_border
    
    # Table 5: Segregasi Entitas
    r_start += 3
    ws_dash.cell(row=r_start, column=1, value='5. Segregasi Sifat Entitas (Barang Fisik vs Jasa & Akun Finansial)').font = Font(name='Calibri', size=12, bold=True, color='1F4E79')
    r_start += 1
    
    dash_h5 = ['No', 'Sifat Entitas', 'Total Item', 'Total Kuantitas', '% Item', 'Rekomendasi Tindakan Logistik & Akuntansi']
    for idx, dh in enumerate(dash_h5, start=1):
        cell = ws_dash.cell(row=r_start, column=idx, value=dh)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = align_center
        
    r_start += 1
    sifat_actions = {
        'Barang Fisik': 'Kelola dalam Master Persediaan Gudang (Stock Card & Opname Rutin)',
        'Jasa (Non-Fisik)': 'Pindahkan ke Master Jasa (Service Item) Accurate; Hapus dari Persediaan',
        'Akun Akuntansi': 'Nonaktifkan dari Master Barang; Input transaksi via Modul Kas/Bank/Jurnal'
    }
    
    row_no = 1
    for sifat_name, sdata in stats_per_sifat.items():
        ws_dash.cell(row=r_start, column=1, value=row_no).alignment = align_center
        ws_dash.cell(row=r_start, column=2, value=sifat_name).font = bold_font
        
        c3 = ws_dash.cell(row=r_start, column=3, value=sdata['total_item'])
        c3.alignment = align_right
        c3.number_format = '#,##0'
        
        c4 = ws_dash.cell(row=r_start, column=4, value=sdata['total_qty'])
        c4.alignment = align_right
        c4.number_format = '#,##0'
        
        pct_sifat = f"{sdata['total_item']/grand_items*100:.1f}%"
        ws_dash.cell(row=r_start, column=5, value=pct_sifat).alignment = align_center
        ws_dash.cell(row=r_start, column=6, value=sifat_actions.get(sifat_name, '-')).alignment = align_left
        
        for c in range(1, 7):
            ws_dash.cell(row=r_start, column=c).border = thin_border
            ws_dash.cell(row=r_start, column=c).font = regular_font
            
        r_start += 1
        row_no += 1

    # Adjust dashboard column widths
    ws_dash.column_dimensions['A'].width = 8
    ws_dash.column_dimensions['B'].width = 42
    ws_dash.column_dimensions['C'].width = 18
    ws_dash.column_dimensions['D'].width = 22
    ws_dash.column_dimensions['E'].width = 24
    ws_dash.column_dimensions['F'].width = 18
    ws_dash.column_dimensions['G'].width = 15
    ws_dash.column_dimensions['H'].width = 16
    ws_dash.column_dimensions['I'].width = 35
    ws_dash.column_dimensions['J'].width = 22

print('Dashboard updated successfully.')

wb.save(OUT_FILE)
wb.save(SRC_FILE)
print(f'Done! Successfully saved to: {OUT_FILE} and {SRC_FILE}')

import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from build_master_dataset import master_crew, vessels, audit_logs

# Create Excel workbook
wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Color palettes (Corporate Maritime Theme)
C_NAVY = "0F2C59"       # Primary Dark Navy
C_HEADER_BG = "1B365D"   # Header Navy Blue
C_ACCENT_BLUE = "0284C7" # Accent Light Blue
C_WHITE = "FFFFFF"
C_LIGHT_GRAY = "F8FAFC"
C_BORDER = "CBD5E1"
C_ALERT_YELLOW = "FEF3C7"
C_ALERT_TEXT = "92400E"
C_VACANT_BG = "FEE2E2"   # Light red/pink
C_VACANT_TEXT = "991B1B" # Dark red
C_CADET_BG = "E0E7FF"    # Light indigo
C_CADET_TEXT = "3730A3"  # Dark indigo
C_ACTIVE_BG = "DCFCE7"   # Light green
C_ACTIVE_TEXT = "166534" # Dark green

font_title = Font(name="Calibri", size=16, bold=True, color=C_NAVY)
font_subtitle = Font(name="Calibri", size=11, italic=True, color="475569")
font_section = Font(name="Calibri", size=13, bold=True, color=C_NAVY)
font_header = Font(name="Calibri", size=11, bold=True, color=C_WHITE)
font_body = Font(name="Calibri", size=10, color="0F172A")
font_body_bold = Font(name="Calibri", size=10, bold=True, color="0F172A")
font_kpi_num = Font(name="Calibri", size=20, bold=True, color=C_NAVY)
font_kpi_label = Font(name="Calibri", size=9, bold=True, color="64748B")

fill_header = PatternFill(start_color=C_HEADER_BG, end_color=C_HEADER_BG, fill_type="solid")
fill_subtotal = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
fill_zebra = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color=C_BORDER),
    right=Side(style='thin', color=C_BORDER),
    top=Side(style='thin', color=C_BORDER),
    bottom=Side(style='thin', color=C_BORDER)
)

double_bottom_border = Border(
    left=Side(style='thin', color=C_BORDER),
    right=Side(style='thin', color=C_BORDER),
    top=Side(style='thin', color=C_BORDER),
    bottom=Side(style='double', color=C_NAVY)
)

# ==========================================
# SHEET 1: RINGKASAN ARMADA & KPI
# ==========================================
ws1 = wb.create_sheet(title="Ringkasan Armada & KPI")
ws1.views.sheetView[0].showGridLines = True

ws1.cell(row=2, column=2, value="PT. CINDARA PRATAMA LINES").font = font_title
ws1.cell(row=3, column=2, value="DATA CREW & MONITORING ARMADA KAPAL - TAHUN 2026 (Per 15 Oktober 2026)").font = font_subtitle
ws1.cell(row=4, column=2, value="Kantor: Jln. Letjend Soeprapto RT.28, No.16, Baru Tengah, Balikpapan 76132 | Telp: 0542-740462").font = Font(name="Calibri", size=9, color="64748B")

# KPI Summary Cards
kpis = [
    ("TOTAL ARMADA", len([v for v in vessels if v['roman'] != 'RELIEF']), C_ACCENT_BLUE, 2),
    ("TOTAL KUOTA KRU", len(master_crew), C_NAVY, 4),
    ("KRU AKTIF", len([c for c in master_crew if c['status_slot'] == 'AKTIF']), "166534", 6),
    ("CADET (MAGANG)", len([c for c in master_crew if 'CADET' in c['status_slot']]), "3730A3", 8),
    ("POSISI LOWONG", len([c for c in master_crew if 'LOWONG' in c['status_slot']]), "991B1B", 10),
]

for title, val, color, col in kpis:
    # Card box row 6 to 7
    ws1.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col+1)
    ws1.merge_cells(start_row=7, start_column=col, end_row=7, end_column=col+1)
    
    cell_lbl = ws1.cell(row=6, column=col, value=title)
    cell_lbl.font = font_kpi_label
    cell_lbl.alignment = Alignment(horizontal="center", vertical="center")
    
    cell_val = ws1.cell(row=7, column=col, value=val)
    cell_val.font = Font(name="Calibri", size=18, bold=True, color=color)
    cell_val.alignment = Alignment(horizontal="center", vertical="center")
    
    for r in range(6, 8):
        for c in range(col, col+2):
            ws1.cell(row=r, column=c).border = thin_border
            ws1.cell(row=r, column=c).fill = PatternFill(start_color="F8FAFC", fill_type="solid")

ws1.cell(row=9, column=2, value="Tabel Rekapitulasi Status Armada Kapal & Kebutuhan Kru").font = font_section

headers_ws1 = [
    "No", "No. Kapal", "Nama Kapal", "Tipe Kapal", "GT", "Mesin (HP)", "Daya (KW/PS)", "No. Telp / Call Kapal",
    "Total Kuota", "Kru Aktif", "Cadet", "Posisi Lowong", "Persentase Terisi"
]

row_idx = 11
for col_idx, h in enumerate(headers_ws1, 2):
    cell = ws1.cell(row=row_idx, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border
ws1.row_dimensions[row_idx].height = 28

vessel_rows = []
for i, v in enumerate(vessels, 1):
    v_crew = [c for c in master_crew if c['vessel_roman'] == v['roman'] and c['vessel_name'] == v['name']]
    total_slot = len(v_crew)
    active_count = len([c for c in v_crew if c['status_slot'] == 'AKTIF'])
    cadet_count = len([c for c in v_crew if 'CADET' in c['status_slot']])
    vacant_count = len([c for c in v_crew if 'LOWONG' in c['status_slot']])
    pct = (active_count + cadet_count) / total_slot if total_slot > 0 else 0
    
    vessel_rows.append({
        'no': i,
        'roman': v['roman'],
        'name': v['name'],
        'type': v['type'],
        'gt': v['gt'],
        'hp': v['hp'],
        'kw': v['kw'] if v['kw'] != '-' else v['ps'],
        'phone': v['phone'],
        'total': total_slot,
        'active': active_count,
        'cadet': cadet_count,
        'vacant': vacant_count,
        'pct': pct
    })

row_idx = 12
for vr in vessel_rows:
    ws1.cell(row=row_idx, column=2, value=vr['no']).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_idx, column=3, value=vr['roman']).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_idx, column=4, value=vr['name']).alignment = Alignment(horizontal="left")
    ws1.cell(row=row_idx, column=5, value=vr['type']).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_idx, column=6, value=vr['gt']).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_idx, column=7, value=vr['hp']).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_idx, column=8, value=vr['kw']).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_idx, column=9, value=vr['phone']).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_idx, column=10, value=vr['total']).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_idx, column=11, value=vr['active']).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_idx, column=12, value=vr['cadet']).alignment = Alignment(horizontal="center")
    
    # Vacant cell highlighting
    cell_vac = ws1.cell(row=row_idx, column=13, value=vr['vacant'])
    cell_vac.alignment = Alignment(horizontal="center")
    if vr['vacant'] > 0:
        cell_vac.font = Font(name="Calibri", size=10, bold=True, color=C_VACANT_TEXT)
        cell_vac.fill = PatternFill(start_color=C_VACANT_BG, fill_type="solid")
    else:
        cell_vac.font = font_body

    cell_pct = ws1.cell(row=row_idx, column=14, value=f"{vr['pct']*100:.0f}%")
    cell_pct.alignment = Alignment(horizontal="center")
    if vr['pct'] == 1.0:
        cell_pct.font = Font(name="Calibri", size=10, color=C_ACTIVE_TEXT)
    else:
        cell_pct.font = Font(name="Calibri", size=10, bold=True, color=C_ALERT_TEXT)

    for c in range(2, 15):
        if c != 13 or vr['vacant'] == 0:
            if row_idx % 2 == 1:
                ws1.cell(row=row_idx, column=c).fill = fill_zebra
        ws1.cell(row=row_idx, column=c).border = thin_border
        if c not in [13, 14]:
            ws1.cell(row=row_idx, column=c).font = font_body
            
    ws1.row_dimensions[row_idx].height = 20
    row_idx += 1

# Total Row
ws1.cell(row=row_idx, column=2, value="TOTAL").alignment = Alignment(horizontal="center")
ws1.merge_cells(start_row=row_idx, start_column=2, end_row=row_idx, end_column=9)
ws1.cell(row=row_idx, column=10, value=f"=SUM(J12:J{row_idx-1})").alignment = Alignment(horizontal="center")
ws1.cell(row=row_idx, column=11, value=f"=SUM(K12:K{row_idx-1})").alignment = Alignment(horizontal="center")
ws1.cell(row=row_idx, column=12, value=f"=SUM(L12:L{row_idx-1})").alignment = Alignment(horizontal="center")
ws1.cell(row=row_idx, column=13, value=f"=SUM(M12:M{row_idx-1})").alignment = Alignment(horizontal="center")
ws1.cell(row=row_idx, column=14, value=f"=(K{row_idx}+L{row_idx})/J{row_idx}").alignment = Alignment(horizontal="center")
ws1.cell(row=row_idx, column=14).number_format = '0.0%'

for c in range(2, 15):
    ws1.cell(row=row_idx, column=c).font = font_body_bold
    ws1.cell(row=row_idx, column=c).fill = fill_subtotal
    ws1.cell(row=row_idx, column=c).border = double_bottom_border
ws1.row_dimensions[row_idx].height = 24


# ==========================================
# SHEET 2: DATA MASTER CREW
# ==========================================
ws2 = wb.create_sheet(title="Data Master Crew")
ws2.views.sheetView[0].showGridLines = True

ws2.cell(row=1, column=1, value="DATA MASTER CREW KAPAL PT. CINDARA PRATAMA LINES - TAHUN 2026").font = font_title
ws2.cell(row=2, column=1, value="Daftar Lengkap Personel On Board, Cadet Magang, dan Posisi Lowong Terstandarisasi").font = font_subtitle

headers_ws2 = [
    "No. ID", "No. Kapal", "Nama Kapal", "Tipe Kapal", "No. Kru", "NIK", "Nama Personel",
    "Status Slot", "Jabatan", "Departemen", "No. Handphone", "Ijazah / Sertifikat",
    "Tahun Ijazah", "Tgl On Board", "Status Penugasan", "Asal Kampus / Sekolah (Cadet)",
    "Riwayat / Catatan Mutasi", "Catatan Audit Data"
]

row_idx = 4
for col_idx, h in enumerate(headers_ws2, 1):
    cell = ws2.cell(row=row_idx, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border
ws2.row_dimensions[row_idx].height = 28

row_idx = 5
for m in master_crew:
    ws2.cell(row=row_idx, column=1, value=m['global_id']).alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=2, value=m['vessel_roman']).alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=3, value=m['vessel_name']).alignment = Alignment(horizontal="left")
    ws2.cell(row=row_idx, column=4, value=m['vessel_type']).alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=5, value=m['vessel_seq']).alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=6, value=m['nik']).alignment = Alignment(horizontal="center")
    
    cell_name = ws2.cell(row=row_idx, column=7, value=m['nama'])
    cell_name.alignment = Alignment(horizontal="left")
    
    cell_status = ws2.cell(row=row_idx, column=8, value=m['status_slot'])
    cell_status.alignment = Alignment(horizontal="center")
    
    # Conditional styling for status slot
    if m['status_slot'] == "LOWONG (VACANT)":
        cell_name.font = Font(name="Calibri", size=10, bold=True, color=C_VACANT_TEXT)
        cell_status.font = Font(name="Calibri", size=10, bold=True, color=C_VACANT_TEXT)
        cell_status.fill = PatternFill(start_color=C_VACANT_BG, fill_type="solid")
    elif m['status_slot'] == "CADET (MAGANG)":
        cell_name.font = Font(name="Calibri", size=10, bold=True, color=C_CADET_TEXT)
        cell_status.font = Font(name="Calibri", size=10, bold=True, color=C_CADET_TEXT)
        cell_status.fill = PatternFill(start_color=C_CADET_BG, fill_type="solid")
    else:
        cell_name.font = font_body
        cell_status.font = Font(name="Calibri", size=10, color=C_ACTIVE_TEXT)

    ws2.cell(row=row_idx, column=9, value=m['jabatan']).alignment = Alignment(horizontal="left")
    ws2.cell(row=row_idx, column=10, value=m['departemen']).alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=11, value=m['hp']).alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=12, value=m['ijazah']).alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=13, value=m['tahun_ijazah']).alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=14, value=m['tgl_onboard']).alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=15, value=m['status_tugas']).alignment = Alignment(horizontal="center")
    ws2.cell(row=row_idx, column=16, value=m['institusi_cadet']).alignment = Alignment(horizontal="left")
    ws2.cell(row=row_idx, column=17, value=m['keterangan']).alignment = Alignment(horizontal="left")
    
    cell_audit = ws2.cell(row=row_idx, column=18, value=m['audit_note'])
    cell_audit.alignment = Alignment(horizontal="left")
    if m['audit_note'] != "-":
        cell_audit.font = Font(name="Calibri", size=9, italic=True, color=C_ALERT_TEXT)
    else:
        cell_audit.font = Font(name="Calibri", size=9, color="94A3B8")

    for c in range(1, 19):
        ws2.cell(row=row_idx, column=c).border = thin_border
        if c not in [7, 8, 18] and m['status_slot'] == "AKTIF":
            ws2.cell(row=row_idx, column=c).font = font_body
            if row_idx % 2 == 1:
                ws2.cell(row=row_idx, column=c).fill = fill_zebra
        elif m['status_slot'] != "AKTIF" and c not in [7, 8, 18]:
            ws2.cell(row=row_idx, column=c).font = font_body

    ws2.row_dimensions[row_idx].height = 20
    row_idx += 1


# ==========================================
# SHEET 3: POSISI LOWONG (VACANT)
# ==========================================
ws3 = wb.create_sheet(title="Posisi Lowong (Vacant)")
ws3.views.sheetView[0].showGridLines = True

ws3.cell(row=1, column=1, value="DAFTAR KEBUTUHAN AWAK KAPAL (POSISI LOWONG / VACANT)").font = Font(name="Calibri", size=16, bold=True, color=C_VACANT_TEXT)
ws3.cell(row=2, column=1, value="Monitoring Kekurangan Formasi Kru Kapal PT. Cindara Pratama Lines untuk Tindak Lanjut Bagian HR & Crewing").font = font_subtitle

headers_ws3 = [
    "No", "No. Kapal", "Nama Kapal", "Tipe Kapal", "Jabatan yang Dibutuhkan", "Departemen",
    "Kualifikasi Ijazah Standar", "No. Telp Kapal", "Rekomendasi Tindakan HR"
]

row_idx = 4
for col_idx, h in enumerate(headers_ws3, 1):
    cell = ws3.cell(row=row_idx, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = PatternFill(start_color="991B1B", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border
ws3.row_dimensions[row_idx].height = 28

vacants_list = [c for c in master_crew if c['status_slot'] == 'LOWONG (VACANT)']
row_idx = 5
for idx, vac in enumerate(vacants_list, 1):
    # find vessel phone
    v_info = next((v for v in vessels if v['roman'] == vac['vessel_roman']), None)
    v_phone = v_info['phone'] if v_info else "-"
    
    ws3.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
    ws3.cell(row=row_idx, column=2, value=vac['vessel_roman']).alignment = Alignment(horizontal="center")
    ws3.cell(row=row_idx, column=3, value=vac['vessel_name']).alignment = Alignment(horizontal="left")
    ws3.cell(row=row_idx, column=4, value=vac['vessel_type']).alignment = Alignment(horizontal="center")
    ws3.cell(row=row_idx, column=5, value=vac['jabatan']).alignment = Alignment(horizontal="left")
    ws3.cell(row=row_idx, column=6, value=vac['departemen']).alignment = Alignment(horizontal="center")
    ws3.cell(row=row_idx, column=7, value=vac['ijazah'] if vac['ijazah'] != '-' else "Sesuai Standar Safe Manning").alignment = Alignment(horizontal="center")
    ws3.cell(row=row_idx, column=8, value=v_phone).alignment = Alignment(horizontal="center")
    
    rekom = "Prioritas Rekrutmen / Mutasi Internal"
    if "Nakhoda" in vac['jabatan'] or "KKM" in vac['jabatan']:
        rekom = "MANDATORI: Segera rekrut / tunjuk Nakhoda/KKM definitif"
    ws3.cell(row=row_idx, column=9, value=rekom).alignment = Alignment(horizontal="left")
    
    for c in range(1, 10):
        ws3.cell(row=row_idx, column=c).border = thin_border
        ws3.cell(row=row_idx, column=c).font = font_body
        if c in [3, 5]:
            ws3.cell(row=row_idx, column=c).font = font_body_bold
        if row_idx % 2 == 1:
            ws3.cell(row=row_idx, column=c).fill = fill_zebra
    ws3.row_dimensions[row_idx].height = 20
    row_idx += 1


# ==========================================
# SHEET 4: MONITORING CADET (MAGANG)
# ==========================================
ws4 = wb.create_sheet(title="Monitoring Cadet")
ws4.views.sheetView[0].showGridLines = True

ws4.cell(row=1, column=1, value="MONITORING CADET (TARUNA PRAKTEK LAUT / MAGANG)").font = Font(name="Calibri", size=16, bold=True, color=C_CADET_TEXT)
ws4.cell(row=2, column=1, value="Daftar Cadet Deck dan Engine yang Sedang Menjalani Praktek Laut (Prala) per 15 Oktober 2026").font = font_subtitle

headers_ws4 = [
    "No", "No. Kapal", "Nama Kapal", "Nama Cadet", "Jabatan", "Departemen",
    "No. Handphone", "Sertifikat", "Tgl On Board", "Asal Kampus / Sekolah Pelayaran", "Catatan"
]

row_idx = 4
for col_idx, h in enumerate(headers_ws4, 1):
    cell = ws4.cell(row=row_idx, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = PatternFill(start_color="3730A3", fill_type="solid")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = thin_border
ws4.row_dimensions[row_idx].height = 28

cadet_list = [c for c in master_crew if 'CADET' in c['status_slot']]
row_idx = 5
for idx, cad in enumerate(cadet_list, 1):
    ws4.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
    ws4.cell(row=row_idx, column=2, value=cad['vessel_roman']).alignment = Alignment(horizontal="center")
    ws4.cell(row=row_idx, column=3, value=cad['vessel_name']).alignment = Alignment(horizontal="left")
    ws4.cell(row=row_idx, column=4, value=cad['nama']).alignment = Alignment(horizontal="left")
    ws4.cell(row=row_idx, column=5, value=cad['jabatan']).alignment = Alignment(horizontal="left")
    ws4.cell(row=row_idx, column=6, value=cad['departemen']).alignment = Alignment(horizontal="center")
    ws4.cell(row=row_idx, column=7, value=cad['hp']).alignment = Alignment(horizontal="center")
    ws4.cell(row=row_idx, column=8, value=cad['ijazah']).alignment = Alignment(horizontal="center")
    ws4.cell(row=row_idx, column=9, value=cad['tgl_onboard']).alignment = Alignment(horizontal="center")
    ws4.cell(row=row_idx, column=10, value=cad['institusi_cadet']).alignment = Alignment(horizontal="left")
    ws4.cell(row=row_idx, column=11, value=cad['audit_note']).alignment = Alignment(horizontal="left")

    for c in range(1, 12):
        ws4.cell(row=row_idx, column=c).border = thin_border
        ws4.cell(row=row_idx, column=c).font = font_body
        if c == 4:
            ws4.cell(row=row_idx, column=c).font = font_body_bold
        if row_idx % 2 == 1:
            ws4.cell(row=row_idx, column=c).fill = fill_zebra
    ws4.row_dimensions[row_idx].height = 20
    row_idx += 1


# ==========================================
# SHEET 5: LOG AUDIT & PERBAIKAN DATA
# ==========================================
ws5 = wb.create_sheet(title="Log Audit & Perbaikan Data")
ws5.views.sheetView[0].showGridLines = True

ws5.cell(row=1, column=1, value="LOG AUDIT & PERBAIKAN STRUKTUR DATA AWAL").font = font_title
ws5.cell(row=2, column=1, value="Dokumentasi Seluruh Inkonsistensi, Typo, dan Kesalahan Format pada Berkas Mentah yang Telah Diperbaiki").font = font_subtitle

headers_ws5 = ["No", "Baris Sumber", "Kategori Masalah", "Deskripsi Temuan & Tindakan Perbaikan"]

row_idx = 4
for col_idx, h in enumerate(headers_ws5, 1):
    cell = ws5.cell(row=row_idx, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = thin_border
ws5.row_dimensions[row_idx].height = 28

row_idx = 5
for idx, log in enumerate(audit_logs, 1):
    ws5.cell(row=row_idx, column=1, value=idx).alignment = Alignment(horizontal="center")
    ws5.cell(row=row_idx, column=2, value=log['baris']).alignment = Alignment(horizontal="center")
    ws5.cell(row=row_idx, column=3, value=log['kategori']).alignment = Alignment(horizontal="center")
    ws5.cell(row=row_idx, column=4, value=log['deskripsi']).alignment = Alignment(horizontal="left")

    for c in range(1, 5):
        ws5.cell(row=row_idx, column=c).border = thin_border
        ws5.cell(row=row_idx, column=c).font = font_body
        if row_idx % 2 == 1:
            ws5.cell(row=row_idx, column=c).fill = fill_zebra
    ws5.row_dimensions[row_idx].height = 20
    row_idx += 1


# Auto-fit all column widths
for ws in wb.worksheets:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # skip merged title rows
            if cell.row < 4 and ws.title != "Ringkasan Armada & KPI":
                continue
            if cell.row < 10 and ws.title == "Ringkasan Armada & KPI":
                continue
            val_str = str(cell.value or "")
            if "\n" in val_str:
                val_str = max(val_str.split("\n"), key=len)
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 11)

# specific tweaks for width
ws1.column_dimensions['A'].width = 4
ws1.column_dimensions['D'].width = 30
ws2.column_dimensions['C'].width = 28
ws2.column_dimensions['G'].width = 28
ws2.column_dimensions['K'].width = 16
ws2.column_dimensions['P'].width = 35
ws2.column_dimensions['Q'].width = 38
ws2.column_dimensions['R'].width = 42

excel_path = r'e:\Dashboard-cpg web\Data_Crew_Kapal_PT_CPL_2026_Clean.xlsx'
wb.save(excel_path)
print(f"Excel successfully saved to: {excel_path}")

# Also export clean CSV
csv_path = r'e:\Dashboard-cpg web\Data_Crew_Kapal_PT_CPL_2026_Clean.csv'
with open(csv_path, 'w', newline='', encoding='utf-8') as f:
    writer = csv.writer(f)
    writer.writerow(headers_ws2)
    for m in master_crew:
        writer.writerow([
            m['global_id'], m['vessel_roman'], m['vessel_name'], m['vessel_type'],
            m['vessel_seq'], m['nik'], m['nama'], m['status_slot'], m['jabatan'],
            m['departemen'], m['hp'], m['ijazah'], m['tahun_ijazah'], m['tgl_onboard'],
            m['status_tugas'], m['institusi_cadet'], m['keterangan'], m['audit_note']
        ])
print(f"CSV successfully saved to: {csv_path}")
